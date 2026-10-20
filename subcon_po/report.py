"""납품리스트/검토사항 xlsx (발주서에 붙여넣기용 + 검토 근거)."""
from __future__ import annotations

from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .engine import Result

HEAD = PatternFill("solid", fgColor="CCFFCC")
WARN = PatternFill("solid", fgColor="FFF2CC")
THIN = Side(style="thin", color="999999")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _table(ws, top: int, headers: list[str], rows: list[list], widths: list[int] | None = None):
    for c, h in enumerate(headers, 1):
        x = ws.cell(top, c, h)
        x.fill, x.font, x.border = HEAD, Font(bold=True), BOX
        x.alignment = Alignment(horizontal="center", vertical="center")
    for i, r in enumerate(rows, 1):
        for c, v in enumerate(r, 1):
            x = ws.cell(top + i, c, v)
            x.border = BOX
            x.alignment = Alignment(vertical="center", wrap_text=isinstance(v, str) and "\n" in v)
    for c, w in enumerate(widths or [], 1):
        ws.column_dimensions[get_column_letter(c)].width = w
    return top + len(rows) + 2


def build_report(res: Result, path: Path) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "수삽_납품리스트"
    ws["A1"] = f"{res.year}년 {res.month}월 생산계획 기준 수삽 일자별 납품모델 (D-1 영업일 / 일수량 +50EA 당김)"
    ws["A1"].font = Font(bold=True, size=13)
    rows = [[d.date_text, d.senjin or None, None, d.note or None, d.model, d.buyer, d.qty,
             d.prod.strftime("%m/%d"), d.plan_row, "O" if d.in_sublist else "미등록"]
            for d in res.deliveries]
    rows += [[r[0], r[1], None, r[2], r[3], r[4], r[5], None, None, "기존 수동행 보존"] for r in res.kept_rows]
    end = _table(ws, 3, ["DATE", "선진행", "", "신흥양산", "MODEL", "BUYER", "Q'TY",
                         "생산계획일", "계획서 행", "SUBLIST"], rows,
                 [10, 8, 3, 28, 30, 22, 8, 11, 9, 14])
    for r in range(4, 4 + len(res.deliveries)):
        if ws.cell(r, 10).value == "미등록":
            for c in range(1, 11):
                ws.cell(r, c).fill = WARN
    ws.cell(end, 1, "※ A~G 열은 임가공 발주서 '수삽' 탭 A~G 와 같은 배치 — 그대로 복사/붙여넣기 가능").font = Font(italic=True)
    ws.freeze_panes = "A4"

    ws2 = wb.create_sheet("검토사항")
    top = 1
    ws2.cell(top, 1, "요약").font = Font(bold=True, size=12)
    top = _table(ws2, top + 1, ["항목", "값"], [
        ["계획 년/월", f"{res.year}년 {res.month}월"],
        ["계획서 시트", res.plan_sheet],
        ["납품 행수", len(res.deliveries)],
        ["납품 수량 합계", sum(d.qty for d in res.deliveries)],
        ["제외 모델 행", len(res.excluded)],
        ["SUBLIST 미등록 모델", len(res.unregistered)],
        ["일정 미정(오더만 있음)", len(res.no_schedule)],
        ["오더≠일자계획 합", len(res.qty_mismatch)],
    ] + [["경고", w] for w in res.warnings], [26, 60, 22, 10, 30])
    ws2.cell(top, 1, "일정 미정 — 일자별 수량이 없어 납품 미반영").font = Font(bold=True)
    top = _table(ws2, top + 1, ["계획서 행", "모델구분", "MODEL", "BUYER", "오더"],
                 [[p.row, p.category, p.model, p.buyer, p.order_qty] for p in res.no_schedule])
    ws2.cell(top, 1, "오더 수량 ≠ 일자계획 합 (일자계획 기준 반영)").font = Font(bold=True)
    top = _table(ws2, top + 1, ["계획서 행", "모델구분", "MODEL", "BUYER", "오더", "일자계획 합"],
                 [[p.row, p.category, p.model, p.buyer, p.order_qty, t] for p, t in res.qty_mismatch])
    ws2.cell(top, 1, "기존 수삽 메모 중 매칭 안 된 행 (선진행/신흥양산 확인 필요)").font = Font(bold=True)
    top = _table(ws2, top + 1, ["DATE", "선진행", "신흥양산", "MODEL", "BUYER", "Q'TY"], res.unmatched_notes)
    ws2.cell(top, 1, "제외 모델").font = Font(bold=True)
    _table(ws2, top + 1, ["계획서 행", "모델구분", "MODEL", "BUYER", "제외 사유", "수량"],
           [[p.row, p.category, p.model, p.buyer, why, sum(p.days.values())] for p, why in res.excluded])
    wb.save(path)
    return path
