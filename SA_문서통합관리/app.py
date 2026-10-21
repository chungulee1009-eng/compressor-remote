"""SA 문서 통합관리 : 승인원 · 설계변경 통보서 PDF 를 한 화면에서 올리고 엑셀 대장으로 정리.

실행 : 실행.bat (또는 python app.py) → 브라우저에서 http://127.0.0.1:8765 자동 열림
저장 : data/설계변경/ , data/승인원/  (대장 엑셀 · 원본 PDF 보관 · 백업)
"""
from __future__ import annotations

import datetime as dt
import os
import re
import shutil
import sys
import threading
import webbrowser
from pathlib import Path

import pdfplumber
from flask import Flask, abort, jsonify, request, send_file, send_from_directory
from openpyxl import load_workbook

from engines import approval, ecn

BASE = Path(__file__).resolve().parent
DATA = Path(os.environ.get("SA_DATA_DIR", BASE / "data")).resolve()
ECN_DIR, APR_DIR = DATA / "설계변경", DATA / "승인원"
UPLOAD_DIR, UNSORTED_DIR = DATA / "_업로드", DATA / "_미분류"
HOST = os.environ.get("SA_HOST", "127.0.0.1")
PORT = int(os.environ.get("SA_PORT", "8765"))

for d in (ECN_DIR, APR_DIR, UPLOAD_DIR, UNSORTED_DIR):
    d.mkdir(parents=True, exist_ok=True)
ecn.set_base(ECN_DIR)
approval.set_base(APR_DIR)

app = Flask(__name__, static_folder=None)
app.config["MAX_CONTENT_LENGTH"] = 300 * 1024 * 1024
LOCK = threading.Lock()   # 대장 엑셀 동시 저장 방지

KIND_LABEL = {"ecn": "설계변경", "approval": "승인원", None: "미분류"}


def norm(s) -> str:
    return re.sub(r"\s+", "", str(s or ""))


def safe_name(name: str) -> str:
    name = os.path.basename(name.replace("\\", "/"))
    name = re.sub(r'[\\/:*?"<>|\x00-\x1f]', "_", name).strip(" .")
    return name or "upload.pdf"


def classify(path: Path):
    """PDF 종류 판별 → ('ecn' | 'approval' | None, 사유)"""
    try:
        with pdfplumber.open(str(path)) as pdf:
            text = norm("".join((p.extract_text() or "") for p in pdf.pages[:4]))
    except Exception as e:
        msg = str(e)
        if "password" in msg.lower() or "encrypt" in msg.lower():
            return None, "암호가 걸린 PDF 입니다. 암호 해제 후 다시 올려주세요."
        return None, "PDF 파일이 손상되었거나 PDF 형식이 아닙니다."
    if len(text) < 10:
        return None, "텍스트가 없는 PDF (스캔 이미지) - 인식 불가"
    if "설계변경통보서" in text:
        return "ecn", ""
    if "承認書" in text or "승인서" in text:
        return "approval", ""
    return None, "승인원 / 설계변경 통보서 양식이 아닙니다."


def classify_xlsx(path: Path):
    """기존 대장 엑셀 가져오기용 판별."""
    try:
        wb = load_workbook(path, read_only=True)
        names = wb.sheetnames
        wb.close()
    except Exception:
        return None
    if "승인원대장" in names:
        return "approval"
    if "상세대장" in names and any(n.startswith("설계변경(") for n in names):
        return "ecn"
    return None


def ledger_info(kind: str) -> dict:
    if kind == "ecn":
        path = ecn.current_ledger()
        count = ecn.ledger_count(path) if path else 0
        name = path.name if path else None
    else:
        path = Path(approval.XLSX) if os.path.exists(approval.XLSX) else None
        count = approval.ledger_count() if path else 0
        name = None
        if path:
            m = dt.datetime.fromtimestamp(path.stat().st_mtime)
            name = f"{m.year}년_SA_승인원 관리대장_{m:%m.%d}.xlsx"
    saved = dt.datetime.fromtimestamp(path.stat().st_mtime).strftime("%Y-%m-%d %H:%M") if path else None
    return {"kind": kind, "label": KIND_LABEL[kind], "exists": bool(path), "count": count,
            "saved": saved, "download_name": name}


def import_ledger(kind: str, path: Path) -> dict:
    """대장이 비어 있을 때만 기존 대장 엑셀을 가져온다 (덮어쓰기 방지)."""
    info = ledger_info(kind)
    if info["count"]:
        return {"kind": "오류", "msg": f"이미 {info['count']}건이 누적된 {info['label']} 대장이 있어 가져오지 않았습니다. "
                                      f"교체하려면 data\\{info['label']} 폴더에서 직접 바꿔주세요."}
    if kind == "ecn":
        for old in ecn.find_ledgers():
            old.unlink()
        wb = load_workbook(path, read_only=True)
        years = sorted(int(m.group(1)) for n in wb.sheetnames if (m := re.fullmatch(r"설계변경\((\d{4})\)", n)))
        wb.close()
        dest = ECN_DIR / ecn.ledger_name(years[-1] if years else dt.date.today().year, dt.date.today())
    else:
        dest = Path(approval.XLSX)
    shutil.copy2(path, dest)
    return {"kind": "등록", "msg": f"기존 {KIND_LABEL[kind]} 대장 가져오기 완료 ({ledger_info(kind)['count']}건). "
                                 f"원본 PDF 링크를 쓰려면 보관 폴더도 data\\{KIND_LABEL[kind]} 에 복사하세요."}


@app.get("/")
def index():
    return send_from_directory(BASE / "web", "index.html")


@app.get("/api/status")
def status():
    return jsonify({"ecn": ledger_info("ecn"), "approval": ledger_info("approval"), "data_dir": str(DATA)})


@app.post("/api/process")
def process():
    files = request.files.getlist("files")
    if not files:
        return jsonify({"error": "업로드된 파일이 없습니다."}), 400
    batch = UPLOAD_DIR / dt.datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    batch.mkdir(parents=True)
    results, groups = [], {"ecn": [], "approval": []}
    errors = []
    with LOCK:
        try:
            saved = []
            for fs in files:
                name = safe_name(fs.filename or "")
                p = batch / name
                n = 2
                while p.exists():
                    p = batch / f"{Path(name).stem}_{n}{Path(name).suffix}"
                    n += 1
                fs.save(p)
                saved.append(p)

            for p in saved:
                ext = p.suffix.lower()
                if ext == ".xlsx":
                    kind = classify_xlsx(p)
                    if kind is None:
                        results.append({"type": None, "file": p.name, "kind": "오류", "key": "",
                                        "msg": "승인원/설계변경 관리대장 엑셀이 아닙니다."})
                    else:
                        r = import_ledger(kind, p)
                        results.append({"type": kind, "file": p.name, "key": "대장 가져오기", **r})
                    continue
                if ext != ".pdf":
                    results.append({"type": None, "file": p.name, "kind": "오류", "key": "",
                                    "msg": "PDF 파일만 등록할 수 있습니다."})
                    continue
                kind, why = classify(p)
                if kind is None:
                    results.append({"type": None, "file": p.name, "kind": "오류", "key": "", "msg": why})
                else:
                    groups[kind].append(p)

            run = {}
            if groups["ecn"]:
                run["ecn"] = ecn.register(groups["ecn"], move=True)
            if groups["approval"]:
                run["approval"] = approval.register([str(p) for p in groups["approval"]], move=True)
            for kind, res in run.items():
                if res.get("error"):
                    errors.append(f"[{KIND_LABEL[kind]}] {res['error']}")
                for it in res["items"]:
                    results.append({"type": kind, **it})
        except Exception as e:  # 예상 못한 오류도 화면에 표시
            errors.append(f"처리 중 오류: {type(e).__name__}: {e}")
        finally:
            # 인식 못 한 PDF 는 _미분류 로 보관, 나머지(중복 등) 임시 파일 삭제
            for r in results:
                src = batch / r["file"]
                if r["kind"] == "오류" and src.exists() and src.suffix.lower() == ".pdf":
                    dst = UNSORTED_DIR / r["file"]
                    if not dst.exists():
                        shutil.move(str(src), str(dst))
            shutil.rmtree(batch, ignore_errors=True)

    order = {"등록": 0, "확인필요": 1, "중복": 2, "오류": 3}
    results.sort(key=lambda r: (order.get(r["kind"], 9), KIND_LABEL[r["type"]]))
    for r in results:
        r["type_label"] = KIND_LABEL[r["type"]]
    return jsonify({"results": results, "errors": errors,
                    "ecn": ledger_info("ecn"), "approval": ledger_info("approval")})


@app.get("/download/<kind>")
def download(kind):
    if kind not in ("ecn", "approval"):
        abort(404)
    info = ledger_info(kind)
    if not info["exists"]:
        abort(404)
    path = ecn.current_ledger() if kind == "ecn" else Path(approval.XLSX)
    return send_file(path, as_attachment=True, download_name=info["download_name"],
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                     max_age=0)


def main():
    url = f"http://{'127.0.0.1' if HOST in ('0.0.0.0', '') else HOST}:{PORT}"
    print("=" * 60)
    print(" SA 문서 통합관리 (승인원 · 설계변경)")
    print(f" 화면 주소 : {url}")
    print(f" 저장 폴더 : {DATA}")
    print(" 이 창을 닫으면 프로그램이 종료됩니다.")
    print("=" * 60)
    import socket
    with socket.socket() as sock:
        sock.settimeout(0.5)
        if sock.connect_ex(("127.0.0.1", PORT)) == 0:     # 이미 실행 중이면 화면만 연다
            print(" 이미 실행 중입니다. 브라우저 화면을 엽니다.")
            webbrowser.open(url)
            return
    if not os.environ.get("SA_NO_BROWSER"):
        threading.Timer(1.2, lambda: webbrowser.open(url)).start()
    try:
        from waitress import serve
        serve(app, host=HOST, port=PORT, threads=4)
    except ImportError:
        app.run(host=HOST, port=PORT, threaded=True)


if __name__ == "__main__":
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    main()
