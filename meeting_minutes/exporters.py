"""회의록 → Excel / Word / PDF, Action Item 현황 → Excel."""
from __future__ import annotations

import re
from datetime import date
from pathlib import Path

from .storage import display_status

NAVY = "1F3864"
STATUS_COLOR = {"완료": "C6EFCE", "진행중": "FFEB9C", "미착수": "F2F2F2", "보류": "D9D9D9", "지연": "FFC7CE"}
STATUS_ICON = {"완료": "🟢", "진행중": "🟡", "미착수": "⚪", "보류": "⏸", "지연": "🔴"}


def safe_filename(s: str) -> str:
    return re.sub(r'[\\/:*?"<>|\s]+', "_", s).strip("_")[:60] or "회의"


def base_name(meeting: dict) -> str:
    return f"{meeting['meeting_date']}_{safe_filename(meeting['title'])}_회의록"


def _sections(m: dict) -> list[tuple[str, list[str]]]:
    mi = m["minutes"]
    return [
        ("Executive Summary", mi.get("summary", [])),
        ("회의 목적", [mi.get("purpose", "")] if mi.get("purpose") else []),
        ("주요 논의사항", [f"[{d['topic']}] {d['content']}" for d in mi.get("discussions", [])]),
        ("결정사항", mi.get("decisions", [])),
        ("문제점", mi.get("issues", [])),
        ("미결사항", mi.get("pending", [])),
        ("차기 회의", [mi["next_meeting"]] if mi.get("next_meeting") else []),
    ]


def _header_rows(m: dict) -> list[tuple[str, str]]:
    mins = int(m.get("duration_sec") or 0) // 60
    return [("회의명", m["title"]), ("일시", f"{m['meeting_date']}" + (f"  (녹음 {mins}분)" if mins else "")),
            ("장소", m.get("location") or "-"), ("참석자", m.get("attendees") or "-")]


def _action_rows(m: dict) -> list[list[str]]:
    rows = []
    for i, a in enumerate(m["actions"], 1):
        st = display_status(a["status"], a["due_date"])
        rows.append([str(i), a["assignee"], a["task"], a["due_date"] or a["due_text"] or "-", a["priority"], st])
    return rows


# ====================================================================== Excel
def export_excel(m: dict, out_dir: str | Path) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    hfill = PatternFill("solid", fgColor=NAVY)
    lfill = PatternFill("solid", fgColor="D9E1F2")
    wrap = Alignment(wrap_text=True, vertical="top")
    font = "맑은 고딕"

    wb = Workbook()
    ws = wb.active
    ws.title = "회의록"
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 95
    ws.merge_cells("A1:B1")
    ws["A1"] = f"회의록 — {m['title']}"
    ws["A1"].font = Font(name=font, size=16, bold=True, color="FFFFFF")
    ws["A1"].fill = hfill
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30
    r = 2
    for k, v in _header_rows(m):
        ws.cell(r, 1, k).font = Font(name=font, bold=True)
        ws.cell(r, 1).fill = lfill
        ws.cell(r, 2, v).font = Font(name=font)
        for c in (1, 2):
            ws.cell(r, c).border = border
        r += 1
    for title, items in _sections(m):
        r += 1
        ws.cell(r, 1, title).font = Font(name=font, bold=True, color="FFFFFF")
        ws.cell(r, 1).fill = hfill
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
        r += 1
        for it in items or ["-"]:
            ws.cell(r, 1, "•").alignment = Alignment(horizontal="right", vertical="top")
            c = ws.cell(r, 2, it)
            c.font = Font(name=font)
            c.alignment = wrap
            ws.row_dimensions[r].height = max(18, 15 * (len(it) // 55 + 1))
            r += 1

    # Action Item 시트
    wa = wb.create_sheet("Action Item")
    heads = ["No", "담당자", "업무", "기한", "우선순위", "상태"]
    widths = [6, 14, 60, 14, 10, 10]
    for i, (h, w) in enumerate(zip(heads, widths), 1):
        c = wa.cell(1, i, h)
        c.font = Font(name=font, bold=True, color="FFFFFF")
        c.fill = hfill
        c.alignment = Alignment(horizontal="center")
        c.border = border
        wa.column_dimensions[c.column_letter].width = w
    for ri, row in enumerate(_action_rows(m), 2):
        for ci, v in enumerate(row, 1):
            c = wa.cell(ri, ci, v)
            c.font = Font(name=font, bold=(ci == 6))
            c.border = border
            c.alignment = wrap if ci == 3 else Alignment(horizontal="center", vertical="top")
        wa.cell(ri, 6).fill = PatternFill("solid", fgColor=STATUS_COLOR.get(row[5], "FFFFFF"))
    wa.freeze_panes = "A2"

    # 전사문 시트
    wt = wb.create_sheet("전사문")
    wt.column_dimensions["A"].width = 12
    wt.column_dimensions["B"].width = 110
    wt.append(["시각", "발언"])
    for ln in (m.get("transcript") or "").splitlines():
        mt = re.match(r"\[(\d{2}:\d{2}:\d{2})\]\s*(.*)", ln)
        wt.append([mt[1], mt[2]] if mt else ["", ln])

    out = Path(out_dir) / f"{base_name(m)}.xlsx"
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


def export_action_tracker(actions: list[dict], out_dir: str | Path, kpi: dict | None = None) -> Path:
    """전체 회의의 Action Item 현황(미결업무 관리표)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Action Item 현황"
    font = "맑은 고딕"
    ws["A1"] = f"Action Item 현황 ({date.today().isoformat()} 기준)"
    ws["A1"].font = Font(name=font, size=14, bold=True)
    if kpi:
        ws["A2"] = (f"전체 {kpi['전체']}건 | 완료 {kpi['완료']} | 진행중 {kpi['진행중']} | 미착수 {kpi['미착수']} | "
                    f"지연 {kpi['지연']} | 완료율 {kpi['완료율']}%")
        ws["A2"].font = Font(name=font, bold=True, color="C00000" if kpi["지연"] else "000000")
    heads = ["회의일", "회의명", "담당자", "업무", "기한", "우선순위", "상태", "메모"]
    widths = [12, 26, 12, 55, 12, 9, 9, 30]
    for i, (h, w) in enumerate(zip(heads, widths), 1):
        c = ws.cell(4, i, h)
        c.font = Font(name=font, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=NAVY)
        ws.column_dimensions[c.column_letter].width = w
    for r, a in enumerate(actions, 5):
        st = display_status(a["status"], a["due_date"])
        vals = [a["meeting_date"], a["meeting_title"], a["assignee"], a["task"], a["due_date"] or a["due_text"],
                a["priority"], st, a.get("note", "")]
        for ci, v in enumerate(vals, 1):
            c = ws.cell(r, ci, v)
            c.font = Font(name=font)
            c.alignment = Alignment(wrap_text=ci in (2, 4, 8), vertical="top")
        ws.cell(r, 7).fill = PatternFill("solid", fgColor=STATUS_COLOR.get(st, "FFFFFF"))
    ws.auto_filter.ref = f"A4:H{max(4, 4 + len(actions))}"
    ws.freeze_panes = "A5"
    out = Path(out_dir) / f"Action_Item_현황_{date.today().isoformat()}.xlsx"
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


# ====================================================================== Word
def export_word(m: dict, out_dir: str | Path) -> Path:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "맑은 고딕"
    st.font.size = Pt(10)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    for s in ("Title", "Heading 1", "Heading 2"):
        doc.styles[s].font.name = "맑은 고딕"
        doc.styles[s].element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")

    def shade(cell, hex_color):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:fill"), hex_color)
        tcPr.append(shd)

    doc.add_heading(f"회의록 — {m['title']}", level=0)
    t = doc.add_table(rows=0, cols=2)
    t.style = "Table Grid"
    for k, v in _header_rows(m):
        row = t.add_row().cells
        row[0].text, row[1].text = k, v
        shade(row[0], "D9E1F2")
        row[0].paragraphs[0].runs[0].bold = True

    for title, items in _sections(m):
        h = doc.add_heading(title, level=1)
        h.runs[0].font.color.rgb = RGBColor(0x1F, 0x38, 0x64)
        for it in items or ["-"]:
            doc.add_paragraph(it, style="List Bullet")
        if title == "문제점":  # Action Item 표는 문제점 다음에
            h = doc.add_heading("Action Item", level=1)
            h.runs[0].font.color.rgb = RGBColor(0x1F, 0x38, 0x64)
            heads = ["No", "담당자", "업무", "기한", "우선순위", "상태"]
            tb = doc.add_table(rows=1, cols=len(heads))
            tb.style = "Table Grid"
            tb.alignment = WD_TABLE_ALIGNMENT.CENTER
            for c, hd in zip(tb.rows[0].cells, heads):
                c.text = hd
                shade(c, NAVY)
                run = c.paragraphs[0].runs[0]
                run.bold = True
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            for r in _action_rows(m):
                cells = tb.add_row().cells
                for c, v in zip(cells, r):
                    c.text = v
                shade(cells[5], STATUS_COLOR.get(r[5], "FFFFFF"))

    out = Path(out_dir) / f"{base_name(m)}.docx"
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)
    return out


# ====================================================================== PDF
_PDF_FONT = None


def _pdf_font() -> str:
    """한글 폰트 등록: 맑은 고딕(Windows) → 나눔고딕 → reportlab 내장 CID 폰트."""
    global _PDF_FONT
    if _PDF_FONT:
        return _PDF_FONT
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    for path in ("C:/Windows/Fonts/malgun.ttf", "C:/Windows/Fonts/gulim.ttc",
                 "/usr/share/fonts/truetype/nanum/NanumGothic.ttf", "/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc",
                 "/Library/Fonts/AppleGothic.ttf", "/System/Library/Fonts/Supplemental/AppleGothic.ttf"):
        if Path(path).exists():
            try:
                pdfmetrics.registerFont(TTFont("KR", path, subfontIndex=0))
                _PDF_FONT = "KR"
                return _PDF_FONT
            except Exception:
                continue
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    pdfmetrics.registerFont(UnicodeCIDFont("HYGothic-Medium"))
    _PDF_FONT = "HYGothic-Medium"
    return _PDF_FONT


def export_pdf(m: dict, out_dir: str | Path) -> Path:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    from xml.sax.saxutils import escape

    f = _pdf_font()
    navy = colors.HexColor("#" + NAVY)
    s_title = ParagraphStyle("t", fontName=f, fontSize=16, leading=22, textColor=navy, spaceAfter=6)
    s_h = ParagraphStyle("h", fontName=f, fontSize=12, leading=16, textColor=navy, spaceBefore=8, spaceAfter=3)
    s_b = ParagraphStyle("b", fontName=f, fontSize=9.5, leading=14, leftIndent=8, bulletIndent=0)
    s_cell = ParagraphStyle("c", fontName=f, fontSize=9, leading=12)
    P = lambda txt, st=s_cell: Paragraph(escape(txt), st)  # noqa: E731

    story = [Paragraph(escape(f"회의록 — {m['title']}"), s_title)]
    hdr = Table([[P(k), P(v)] for k, v in _header_rows(m)], colWidths=[30 * mm, 145 * mm])
    hdr.setStyle(TableStyle([("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#D9E1F2")),
                             ("GRID", (0, 0), (-1, -1), 0.4, colors.grey), ("VALIGN", (0, 0), (-1, -1), "TOP")]))
    story.append(hdr)
    for title, items in _sections(m):
        story.append(Paragraph(escape(title), s_h))
        for it in items or ["-"]:
            story.append(Paragraph(escape(it), s_b, bulletText="•"))
        if title == "문제점":
            story.append(Paragraph("Action Item", s_h))
            s_white = ParagraphStyle("w", parent=s_cell, textColor=colors.white)
            arows = _action_rows(m)
            rows = [[P(h, s_white) for h in ["No", "담당자", "업무", "기한", "우선", "상태"]]]
            rows += [[P(v) for v in r] for r in arows]
            style = [("BACKGROUND", (0, 0), (-1, 0), navy),
                     ("GRID", (0, 0), (-1, -1), 0.4, colors.grey), ("VALIGN", (0, 0), (-1, -1), "TOP")]
            for i, r in enumerate(arows, 1):
                style.append(("BACKGROUND", (5, i), (5, i), colors.HexColor("#" + STATUS_COLOR.get(r[5], "FFFFFF"))))
            tb = Table(rows, colWidths=[10 * mm, 24 * mm, 85 * mm, 24 * mm, 14 * mm, 18 * mm], repeatRows=1)
            tb.setStyle(TableStyle(style))
            story.append(tb)
        story.append(Spacer(1, 2))

    out = Path(out_dir) / f"{base_name(m)}.pdf"
    out.parent.mkdir(parents=True, exist_ok=True)
    SimpleDocTemplate(str(out), pagesize=A4, leftMargin=17 * mm, rightMargin=17 * mm, topMargin=15 * mm,
                      bottomMargin=15 * mm, title=f"회의록 {m['title']}").build(story)
    return out
