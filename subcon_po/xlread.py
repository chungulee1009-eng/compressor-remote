"""xlsx/xlsm/xlsb 공통 읽기: 시트 → 2차원 값 리스트(0-based)."""
from __future__ import annotations

from pathlib import Path

import openpyxl
from openpyxl.reader import workbook as _wbr

# 일부 발주서의 '인쇄 제목' 정의명이 비어 있어 openpyxl 이 예외를 냄 → 무시
_orig_assign = _wbr.WorkbookParser.assign_names


def _safe_assign(self):
    try:
        _orig_assign(self)
    except TypeError:
        pass


_wbr.WorkbookParser.assign_names = _safe_assign


def sheet_names(path) -> list[str]:
    path = Path(path)
    if path.suffix.lower() == ".xlsb":
        from pyxlsb import open_workbook
        with open_workbook(str(path)) as wb:
            return list(wb.sheets)
    wb = openpyxl.load_workbook(path, read_only=True)
    try:
        return [ws.title for ws in wb.worksheets if ws.sheet_state == "visible"] + \
               [ws.title for ws in wb.worksheets if ws.sheet_state != "visible"]
    finally:
        wb.close()


def read_sheets(path, names: list[str] | None = None) -> dict[str, list[list]]:
    """names=None 이면 전체 시트. 값은 캐시된 계산값(수식 결과)."""
    path = Path(path)
    out: dict[str, list[list]] = {}
    if path.suffix.lower() == ".xlsb":
        from pyxlsb import open_workbook
        with open_workbook(str(path)) as wb:
            for name in wb.sheets:
                if names is not None and name not in names:
                    continue
                grid: list[list] = []
                with wb.get_sheet(name) as sh:
                    for row in sh.rows(sparse=True):
                        if not row:
                            continue
                        r = row[0].r
                        while len(grid) <= r:
                            grid.append([])
                        vals = grid[r]
                        for c in row:
                            if c.v is None:
                                continue
                            while len(vals) <= c.c:
                                vals.append(None)
                            vals[c.c] = c.v
                out[name] = grid
        return out
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        for ws in wb.worksheets:
            if names is not None and ws.title not in names:
                continue
            out[ws.title] = [list(r) for r in ws.iter_rows(values_only=True)]
    finally:
        wb.close()
    return out


def cell(grid: list[list], r: int, c: int):
    if r < len(grid) and c < len(grid[r]):
        return grid[r][c]
    return None
