"""임가공 발주서 '수삽' 탭 A~G 기입 후 저장.

백엔드 우선순위
  1) Excel(COM, Windows + Excel 설치)   : .xlsb 그대로 저장, 서식/매크로/수식 100% 보존
  2) LibreOffice(UNO, soffice 설치)     : .xlsb 입력 → .xlsx 저장
  3) openpyxl                           : .xlsx/.xlsm 입력만 (도형/이미지 일부 손실 가능)
"""
from __future__ import annotations

import json
import os
import shutil
import subprocess
import sys
from pathlib import Path

from .engine import SUSAP

# 기입 열(1-based): A DATE, B 선진행, D 신흥양산, E MODEL, F BUYER, G Q'TY  (C 미사용)
WRITE_COLS = [1, 2, 4, 5, 6, 7]


def _excel_com(src: Path, dst: Path, start: int, end: int, rows: list[list]) -> None:
    import pythoncom
    import win32com.client

    fmt = {".xlsb": 50, ".xlsx": 51, ".xlsm": 52, ".xls": 56}[dst.suffix.lower()]
    pythoncom.CoInitialize()
    xl = win32com.client.DispatchEx("Excel.Application")
    xl.Visible = False
    xl.DisplayAlerts = False
    xl.ScreenUpdating = False
    wb = None
    try:
        wb = xl.Workbooks.Open(str(src.resolve()), UpdateLinks=0)
        ws = wb.Worksheets(SUSAP)
        ws.Range(f"A{start}:B{end}").ClearContents()
        ws.Range(f"D{start}:G{end}").ClearContents()
        if rows:
            last = start + len(rows) - 1
            ws.Range(f"A{start}:B{last}").Value = tuple((r[0], r[1]) for r in rows)
            ws.Range(f"D{start}:G{last}").Value = tuple((r[2], r[3], r[4], r[5]) for r in rows)
        wb.SaveAs(str(dst.resolve()), FileFormat=fmt)
    finally:
        if wb is not None:
            wb.Close(SaveChanges=False)
        xl.Quit()
        pythoncom.CoUninitialize()


def _find_soffice() -> str | None:
    for name in ("soffice", "libreoffice"):
        p = shutil.which(name)
        if p:
            return p
    for p in (r"C:\Program Files\LibreOffice\program\soffice.exe",
              r"C:\Program Files (x86)\LibreOffice\program\soffice.exe"):
        if os.path.exists(p):
            return p
    return None


def _uno_python(soffice: str) -> str | None:
    """uno 모듈을 import 할 수 있는 파이썬 (현재 인터프리터 → LO 동봉 파이썬)."""
    try:
        import uno  # noqa: F401
        return sys.executable
    except ImportError:
        pass
    cand = Path(soffice).with_name("python.exe" if os.name == "nt" else "python")
    return str(cand) if cand.exists() else None


def _libreoffice(src: Path, dst: Path, start: int, end: int, rows: list[list]) -> None:
    soffice = _find_soffice()
    py = _uno_python(soffice) if soffice else None
    if not py:
        raise RuntimeError("LibreOffice(UNO) 를 찾지 못했습니다.")
    job = dst.with_suffix(".job.json")
    job.write_text(json.dumps({"soffice": soffice, "src": str(src.resolve()), "dst": str(dst.resolve()),
                               "sheet": SUSAP, "start": start, "end": end, "cols": WRITE_COLS,
                               "rows": rows}, ensure_ascii=False), encoding="utf-8")
    try:
        script = Path(__file__).with_name("lo_write.py")
        p = subprocess.run([py, str(script), str(job)], capture_output=True, text=True, timeout=600)
        if p.returncode != 0 or not dst.exists():
            raise RuntimeError(f"LibreOffice 저장 실패: {(p.stderr or p.stdout)[-800:]}")
    finally:
        job.unlink(missing_ok=True)


def _openpyxl(src: Path, dst: Path, start: int, end: int, rows: list[list]) -> None:
    from . import xlread  # noqa: F401  (정의명 예외 패치 적용)
    import openpyxl

    wb = openpyxl.load_workbook(src, keep_vba=src.suffix.lower() == ".xlsm")
    ws = wb[SUSAP]
    for r in range(start, end + 1):
        for c in WRITE_COLS:
            ws.cell(r, c).value = None
    for i, vals in enumerate(rows):
        for c, v in zip(WRITE_COLS, vals):
            ws.cell(start + i, c).value = v
    wb.save(dst)


def available_backends(src_suffix: str) -> list[str]:
    out = []
    if os.name == "nt":
        try:
            import win32com.client  # noqa: F401
            out.append("excel")
        except ImportError:
            pass
    if _find_soffice():
        out.append("libreoffice")
    if src_suffix.lower() in (".xlsx", ".xlsm"):
        out.append("openpyxl")
    return out


def write_order(src, out_dir, start: int, end: int, rows: list[list], stem: str) -> tuple[Path, str]:
    """수삽 탭을 갱신한 발주서를 out_dir 에 저장. (경로, 사용 백엔드) 반환."""
    src, out_dir = Path(src), Path(out_dir)
    errors = []
    for be in available_backends(src.suffix):
        if be == "excel":
            dst = out_dir / f"{stem}{src.suffix}"
            fn = _excel_com
        elif be == "libreoffice":
            dst = out_dir / f"{stem}.xlsx" if src.suffix.lower() == ".xlsb" else out_dir / f"{stem}{src.suffix}"
            fn = _libreoffice
        else:
            dst = out_dir / f"{stem}{src.suffix}"
            fn = _openpyxl
        try:
            dst.unlink(missing_ok=True)
            fn(src, dst, start, end, rows)
            return dst, be
        except Exception as e:  # 다음 백엔드로
            errors.append(f"{be}: {e}")
    raise RuntimeError("발주서 저장 불가 — " + (" / ".join(errors) if errors else
                       ".xlsb 는 Excel(Windows) 또는 LibreOffice 가 필요합니다. 납품리스트 파일을 붙여넣어 사용하세요."))
