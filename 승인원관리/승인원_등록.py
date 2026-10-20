# -*- coding: utf-8 -*-
"""
승인원 PDF -> 엑셀 누적 관리대장 자동 등록

사용법
  1) 승인원_수신함 폴더에 PDF 를 넣고 실행 (인자 없음)
     -> 등록 후 PDF 는 승인원_보관/<연도>/ 로 이동
  2) python 승인원_등록.py a.pdf b.pdf ...  (또는 .bat 에 PDF 드래그&드롭)
     -> 등록 후 PDF 를 승인원_보관/<연도>/ 로 복사

결과: 승인원_관리대장.xlsx
  - 승인원대장   : 승인원 1건 = 1행 (누적, 비고 열은 자유 기입)
  - 검토항목     : 승인서 검토/시험 항목별 판정
  - 승인이력     : 승인원 내 버전 변경 이력
  - 품번별최신   : 품번별 최신 승인 버전 (실행 시 자동 재생성)
  - 현황요약     : 업체/모델/승인구분/월별 집계 (COUNTIFS 수식)
"""
import datetime as dt
import hashlib
import os
import re
import shutil
import sys
from collections import OrderedDict

try:
    import pdfplumber
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError as e:  # pragma: no cover
    print(f"[오류] 필요한 패키지가 없습니다: {e.name}")
    print("       pip install -r requirements.txt 실행 후 다시 시도하세요.")
    sys.exit(1)

BASE = os.path.dirname(os.path.abspath(__file__))
XLSX = os.path.join(BASE, "승인원_관리대장.xlsx")
INBOX = os.path.join(BASE, "승인원_수신함")
ARCHIVE = os.path.join(BASE, "승인원_보관")
ERROR_DIR = os.path.join(INBOX, "오류")
DUP_DIR = os.path.join(INBOX, "중복")
BACKUP_DIR = os.path.join(BASE, "백업")
BACKUP_KEEP = 20

FONT = "맑은 고딕"

# ---------------------------------------------------------------- 대장 컬럼
LEDGER_COLS = [
    ("No", 6), ("등록일시", 17), ("발행No", 22), ("승인일자", 12),
    ("승인구분", 14), ("버전", 7), ("품번", 15), ("품목이름", 34),
    ("업체 P/N", 20), ("적용모델", 12), ("적용시점", 9), ("제작업체", 15),
    ("공급업체", 15), ("검토수량", 9), ("포장정보", 10),
    ("외관", 8), ("치수", 8), ("성능", 8), ("기타", 8), ("최종판정", 9),
    ("작성부서", 12), ("작성자", 12), ("검토자", 9), ("심사자", 9), ("승인자", 9),
    ("배포선", 22), ("업체담당자", 12), ("업체연락처", 15), ("업체이메일", 24),
    ("이전버전", 8), ("변경내역", 30), ("특기사항", 50),
    ("페이지수", 8), ("원본파일", 40), ("파싱상태", 18), ("파일해시", 14), ("비고", 30),
]
LEDGER_HEAD = [c for c, _ in LEDGER_COLS]
COL = {name: i + 1 for i, name in enumerate(LEDGER_HEAD)}

ITEM_COLS = [("발행No", 22), ("품번", 15), ("버전", 7), ("NO", 5), ("구분", 9),
             ("검토 및 시험사항", 50), ("판정", 8)]
HIST_COLS = [("발행No", 22), ("품번", 15), ("이력No", 7), ("승인일자", 12),
             ("변경 전 버전", 11), ("변경 후 버전", 11), ("변경 전 내용", 30),
             ("변경 후 내용", 30), ("구형재고", 10), ("변경시점", 12)]

# ---------------------------------------------------------------- 서식
C_HEAD = PatternFill("solid", fgColor="1F3864")
C_SUB = PatternFill("solid", fgColor="D9E1F2")
C_NOTE = PatternFill("solid", fgColor="FFF2CC")
C_BAD = PatternFill("solid", fgColor="F8CBAD")
F_HEAD = Font(name=FONT, bold=True, color="FFFFFF", size=10)
F_BODY = Font(name=FONT, size=10)
F_BOLD = Font(name=FONT, size=10, bold=True)
F_LINK = Font(name=FONT, size=10, color="0563C1", underline="single")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)


def nz(s):
    """공백/개행 제거한 비교용 문자열"""
    return re.sub(r"\s+", "", s or "")


def clean(s):
    return re.sub(r"[ \t]+", " ", (s or "").replace("\r", "")).strip()


def judge(s):
    """판정 문구만 인식. 일반 문장(예: PACKING 의 'NG')을 판정으로 오인하지 않도록 완전일치 위주"""
    t = nz(s).upper()
    if not t:
        return ""
    if "불합격" in t or "不合格" in t or "부적합" in t or t in ("NG", "N.G", "FAIL", "X"):
        return "불합격"
    if "합격" in t or "合格" in t or t in ("OK", "O.K", "PASS", "적합", "양호", "○"):
        return "합격"
    return ""


def parse_date(s):
    m = re.search(r"(20\d{2})\s*[년.\-/]\s*(\d{1,2})\s*[월.\-/]\s*(\d{1,2})", s or "")
    if m:
        try:
            return dt.date(int(m[1]), int(m[2]), int(m[3]))
        except ValueError:
            return None
    return None


# ---------------------------------------------------------------- PDF 파싱
LABELS = {
    "작성부서": "작성부서", "작성자": "작성자", "발행No.": "발행No", "발행No": "발행No",
    "품번": "품번", "적용모델": "적용모델", "적용시점": "적용시점", "품목이름": "품목이름",
    "검토수량": "검토수량", "제작업체": "제작업체", "포장정보": "포장정보",
    "공급업체": "공급업체", "주문정보": "업체 P/N",
}


def _next_value(row, i):
    for v in row[i + 1:]:
        if v is not None and clean(v):
            return v
    return ""


def parse_approval_page(page, rec, warn):
    text = page.extract_text() or ""
    m = re.search(r"\(\s*([^)]+?)\s*\)\s*承認書", text)
    rec["_서식"] = f"({clean(m[1])})承認書" if m else "承認書"
    d = parse_date(text)
    rec["승인일자"] = d

    rows = [r for t in page.extract_tables() for r in t]
    in_items, cur_no, cur_cat, jcol = False, "", "", None
    sign_cols = None
    dist_head = None
    items = []
    for ri, row in enumerate(rows):
        cells = [c if c is None else c.replace("\r", "") for c in row]
        keys = [nz(c) for c in cells]

        # 결재란: 헤더(작성/검토/심사/승인) -> 다음 행에 이름+일자
        if sign_cols is None and "작성" in keys and "승인" in keys and "검토" in keys:
            sign_cols = {k: keys.index(k) for k in ("작성", "검토", "심사", "승인") if k in keys}
            if ri + 1 < len(rows):
                nxt = rows[ri + 1]
                for k, ci in sign_cols.items():
                    v = nxt[ci] if ci < len(nxt) else None
                    rec[{"작성": "_작성결재", "검토": "검토자", "심사": "심사자", "승인": "승인자"}[k]] = \
                        clean((v or "").split("\n")[0])

        # 라벨 -> 값
        for ci, k in enumerate(keys):
            if k in LABELS and LABELS[k] not in rec:
                rec[LABELS[k]] = clean(_next_value(cells, ci).replace("\n", "" if k.startswith("발행") else " "))
            if k.startswith("검토및이유이력등"):
                rec["_검토이유"] = clean(cells[ci].split("\n", 1)[1]) if "\n" in cells[ci] else \
                    clean(_next_value(cells, ci))

        # 검토/시험 항목
        if keys and keys[0] == "NO" and any("검토및시험사항" in k for k in keys):
            in_items = True
            jcol = next((ci for ci, k in enumerate(keys) if k == "판정"), None)
            continue
        if in_items:
            if keys and keys[0].startswith("특기사항"):
                in_items = False
            else:
                vals = [clean(c) for c in cells if c is not None]
                body = [v for v in vals if v]
                if not body:
                    continue
                no = clean(cells[0]) if cells[0] else ""
                cat = clean(cells[1]) if len(cells) > 1 and cells[1] else ""
                cur_no = no or cur_no
                cur_cat = cat or cur_cat
                raw = ""
                if jcol is not None:
                    raw = next((clean(c) for c in cells[jcol:] if c is not None and clean(c)), "")
                elif len(body) >= 2:
                    raw = body[-1]
                res = judge(raw)
                desc = [v for v in body if v not in (no, cat)]
                if raw and desc and desc[-1] == raw:
                    desc = desc[:-1]
                items.append({"NO": cur_no, "구분": cur_cat,
                              "검토 및 시험사항": clean(" ".join(desc)).lstrip("- ").strip(),
                              "판정": res})

        if keys and keys[0].startswith("특기사항"):
            notes = cells[0].split("\n")[1:]
            rec["특기사항"] = "\n".join(clean(n) for n in notes if clean(n))

        # 배포선 / 최종판정
        if keys and keys[0] == "배포선":
            dist_head = cells
            for ci, k in enumerate(keys):
                if k.endswith("승인") and k != "승인":
                    rec["승인구분"] = rec.get("승인구분") or k
                if judge(k) and k.upper() not in ("O", "○", "X"):
                    rec["최종판정"] = judge(k)
        if keys and keys[0] == "승인원" and dist_head:
            depts = []
            for ci, k in enumerate(keys):
                if k.upper() in ("O", "○", "●", "V", "√") and ci < len(dist_head):
                    depts.append(clean(dist_head[ci]))
            rec["배포선"] = ", ".join(d for d in depts if d)

    rec["_items"] = items

    # 표 인식 실패 대비: 텍스트 정규식 보완
    if not rec.get("품번"):
        m = re.search(r"품\s*번\s+(\S+)", text)
        if m:
            rec["품번"] = m[1]
            warn.append("품번 텍스트추출")
    if not rec.get("발행No"):
        m = re.search(r"([A-Z]{1,4}\s*승인-\d{2}-\d{3,5}(?:-\S+)?)", text.replace("\n", ""))
        if m:
            rec["발행No"] = nz(m[1])
            warn.append("발행No 텍스트추출")


def parse_history(pages, rec):
    """승인이력 표는 제목과 다른 페이지에 있을 수 있어 앞쪽 페이지 전체에서 행 패턴으로 탐색"""
    hist, seen = [], set()
    for page in pages:
        for t in page.extract_tables():
            for row in t:
                cells = [clean((c or "").replace("\n", " ")) for c in row]
                if len(cells) >= 8 and re.fullmatch(r"\d+", cells[0]) and parse_date(cells[1]):
                    key = (cells[0], cells[1])
                    if key in seen:
                        continue
                    seen.add(key)
                    hist.append({"이력No": int(cells[0]), "승인일자": parse_date(cells[1]),
                                 "변경 전 버전": cells[2], "변경 후 버전": cells[3],
                                 "변경 전 내용": cells[4], "변경 후 내용": cells[5],
                                 "구형재고": cells[6], "변경시점": cells[7]})
        text = page.extract_text() or ""
        m = re.search(r"변경\s*상세\s*내역\s*\n(.+)", text, re.S)
        if m and "변경내역" not in rec:
            rec["변경내역"] = "\n".join(clean(x).lstrip("- ").strip() for x in m[1].split("\n") if clean(x))
    rec["_hist"] = hist


def special_notes(page):
    """특기사항: 좌측 '*' 항목 + 우측 별도 메모를 분리 (줄 좌표 기준)"""
    lines = page.extract_text_lines()
    start = next((l for l in lines if nz(l["text"]).startswith("특기사항")), None)
    if not start:
        return None
    end = next((l for l in lines if l["top"] > start["top"] and nz(l["text"]).startswith("배포선")), None)
    left, right = [], []
    for l in lines:
        if l["top"] <= start["top"] + 1 or (end and l["top"] >= end["top"] - 1):
            continue
        t = clean(l["text"])
        if l["x0"] > page.width * 0.5:
            right.append(t)
        elif t.startswith("*") or not left:
            left.append(t)
        else:
            left[-1] += " " + t  # 줄바꿈된 항목 이어붙임
    if right:
        left.append("※ " + " ".join(right))
    return "\n".join(left)


def parse_vendor_page(page, rec):
    """시험성적서 업체정보 표: 공급자 담당자 우선, 없으면 제조자"""
    info, sect = {"제조자": {}, "공급자": {}}, None
    fields = {"이름": "업체담당자", "전화번호": "전화", "핸드폰번호": "HP",
              "이메일": "업체이메일", "이메일주소": "업체이메일"}
    for t in page.extract_tables():
        for row in t:
            keys = [nz(c) for c in row]
            if keys and keys[0] in info:
                sect = keys[0]
            if not sect:
                continue
            for ci, k in enumerate(keys):
                if k in fields:
                    v = clean(_next_value(row, ci))
                    if v and v != "-":
                        info[sect].setdefault(fields[k], v)
    for sect in ("공급자", "제조자"):
        d = info[sect]
        rec.setdefault("업체담당자", d.get("업체담당자"))
        tel = " / ".join(x for x in (d.get("전화"), d.get("HP") and "HP " + d["HP"]) if x)
        if tel and not rec.get("업체연락처"):
            rec["업체연락처"] = tel
        if d.get("업체이메일") and not rec.get("업체이메일"):
            rec["업체이메일"] = d["업체이메일"]
    if not rec.get("업체담당자"):
        rec.pop("업체담당자", None)


def parse_pdf(path):
    rec, warn = {}, []
    with pdfplumber.open(path) as pdf:
        rec["페이지수"] = len(pdf.pages)
        front = pdf.pages[:6]
        texts = [(p.extract_text() or "") for p in front]
        idx = next((i for i, t in enumerate(texts) if "承認書" in t or "승인서" in nz(t)), 0)
        parse_approval_page(pdf.pages[idx], rec, warn)
        notes = special_notes(pdf.pages[idx])
        if notes:
            rec["특기사항"] = notes
        parse_history(front, rec)
        vi = next((i for i, t in enumerate(texts) if "업체정보" in nz(t)), None)
        if vi is not None:
            parse_vendor_page(pdf.pages[vi], rec)

    # 승인구분 / 버전
    reason = rec.get("_검토이유", "")
    m = re.search(r"^(.*?)\s*\(\s*(V[\d.]+)\s*\)", reason)
    if m:
        rec["승인구분"] = clean(m[1]) or rec.get("승인구분", "")
        rec["버전"] = m[2]
    elif reason:
        rec["승인구분"] = reason
    if not rec.get("버전"):
        m = re.search(r"\b(V\d+(?:\.\d+)*)\s*으?로\s*관리", rec.get("특기사항", ""))
        if m:
            rec["버전"] = m[1]
        elif rec["_hist"]:
            rec["버전"] = rec["_hist"][-1]["변경 후 버전"]
    if rec["_hist"]:
        last = rec["_hist"][-1]
        rec["이전버전"] = last["변경 전 버전"] or "-"
        rec.setdefault("변경내역", last["변경 후 내용"])
        rec["승인일자"] = rec.get("승인일자") or last["승인일자"]

    # 작성자: 결재란 우선, 없으면 작성자 라벨
    rec["작성자"] = rec.get("작성자") or rec.get("_작성결재", "")

    # 구분별 판정 집계
    for cat in ("외관", "치수", "성능", "기타"):
        js = [it["판정"] for it in rec["_items"] if nz(it["구분"]) == cat and it["판정"]]
        rec[cat] = "불합격" if "불합격" in js else ("합격" if js else "")
    if not rec.get("최종판정"):
        js = [it["판정"] for it in rec["_items"] if it["판정"]]
        rec["최종판정"] = "불합격" if "불합격" in js else ("합격" if js else "")
        if js:
            warn.append("최종판정=항목집계(승인도장 미인식)")
    bad = [f"{it['구분']}:{it['검토 및 시험사항'][:20]}" for it in rec["_items"] if it["판정"] == "불합격"]
    if rec.get("최종판정") == "불합격" or bad:
        warn.append("불합격근거=" + ("; ".join(bad) if bad else "승인도장 불합격"))

    missing = [k for k in ("발행No", "품번", "승인일자", "버전", "최종판정") if not rec.get(k)]
    if missing:
        warn.append("누락:" + "/".join(missing))
    rec["파싱상태"] = "정상" if not warn else "확인필요(" + ", ".join(warn) + ")"
    return rec


# ---------------------------------------------------------------- 엑셀
def style_header(ws, cols, row=1):
    for i, (name, width) in enumerate(cols, 1):
        c = ws.cell(row=row, column=i, value=name)
        c.font, c.fill, c.alignment, c.border = F_HEAD, C_HEAD, CENTER, BORDER
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[row].height = 24
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def style_row(ws, r, ncol, wrap_cols=()):
    for ci in range(1, ncol + 1):
        c = ws.cell(row=r, column=ci)
        if c.font != F_LINK:
            c.font = F_BODY
        c.border = BORDER
        c.alignment = LEFT if ci in wrap_cols else CENTER


def open_workbook():
    if os.path.exists(XLSX):
        return load_workbook(XLSX)
    wb = Workbook()
    ws = wb.active
    ws.title = "승인원대장"
    style_header(ws, LEDGER_COLS)
    for title, cols in (("검토항목", ITEM_COLS), ("승인이력", HIST_COLS)):
        style_header(wb.create_sheet(title), cols)
    wb.create_sheet("품번별최신")
    wb.create_sheet("현황요약", 0)
    return wb


def ledger_rows(ws):
    out = []
    for r in range(2, ws.max_row + 1):
        if ws.cell(row=r, column=COL["발행No"]).value or ws.cell(row=r, column=COL["품번"]).value:
            out.append(r)
    return out


def archive_name(rec, src):
    year = str((rec.get("승인일자") or dt.date.today()).year)
    stem = "_".join(x for x in (rec.get("발행No"), rec.get("품번"), rec.get("버전")) if x)
    stem = re.sub(r'[\\/:*?"<>|\s]+', "_", stem) or os.path.splitext(os.path.basename(src))[0]
    rel = os.path.join("승인원_보관", year, stem + ".pdf")
    n = 2
    while os.path.exists(os.path.join(BASE, rel)):
        rel = os.path.join("승인원_보관", year, f"{stem}({n}).pdf")
        n += 1
    return rel


KEEP_ON_REPARSE = ("No", "등록일시", "원본파일", "파일해시", "비고")


def write_ledger_row(ws, r, rec):
    for name in LEDGER_HEAD:
        if name in KEEP_ON_REPARSE and name not in rec:
            continue
        v = rec.get(name, "")
        ws.cell(row=r, column=COL[name], value=v if v != "" else None)
    for name in ("최종판정", "파싱상태"):
        ws.cell(row=r, column=COL[name]).fill = PatternFill(fill_type=None)


def append_ledger(ws, rec, rel_path, sha):
    r = max(ledger_rows(ws) or [1]) + 1
    vals = dict(rec)
    vals.update({"No": r - 1, "등록일시": dt.datetime.now().replace(microsecond=0),
                 "원본파일": os.path.basename(rel_path), "파일해시": sha[:12]})
    write_ledger_row(ws, r, vals)
    finish_ledger_row(ws, r, rec)
    link = ws.cell(row=r, column=COL["원본파일"])
    link.hyperlink = rel_path.replace(os.sep, "/")
    link.font = F_LINK
    return r


def finish_ledger_row(ws, r, rec):
    style_row(ws, r, len(LEDGER_HEAD),
              wrap_cols=(COL["품목이름"], COL["변경내역"], COL["특기사항"], COL["비고"], COL["원본파일"],
                         COL["파싱상태"]))
    ws.cell(row=r, column=COL["승인일자"]).number_format = "yyyy-mm-dd"
    ws.cell(row=r, column=COL["등록일시"]).number_format = "yyyy-mm-dd hh:mm"
    if rec.get("최종판정") == "불합격":
        ws.cell(row=r, column=COL["최종판정"]).fill = C_BAD
    if rec.get("파싱상태") != "정상":
        ws.cell(row=r, column=COL["파싱상태"]).fill = C_NOTE
    ws.auto_filter.ref = f"A1:{get_column_letter(len(LEDGER_HEAD))}{max(ledger_rows(ws) or [r])}"


def append_simple(ws, cols, rows, wrap_names=()):
    wrap = tuple(i + 1 for i, (n, _) in enumerate(cols) if n in wrap_names)
    r = ws.max_row
    while r > 1 and ws.cell(row=r, column=1).value is None:
        r -= 1
    for row in rows:
        r += 1
        for i, (name, _) in enumerate(cols, 1):
            ws.cell(row=r, column=i, value=row.get(name))
            if isinstance(row.get(name), dt.date):
                ws.cell(row=r, column=i).number_format = "yyyy-mm-dd"
        style_row(ws, r, len(cols), wrap)
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{max(r, 1)}"


def _ver_key(v):
    return tuple(int(x) for x in re.findall(r"\d+", v or "")) or (-1,)


def rebuild_latest(wb):
    ws_l = wb["승인원대장"]
    latest = OrderedDict()
    counts = {}
    for r in ledger_rows(ws_l):
        g = lambda n: ws_l.cell(row=r, column=COL[n]).value
        pn = g("품번") or "(품번없음)"
        counts[pn] = counts.get(pn, 0) + 1
        d = g("승인일자")
        d = d.date() if isinstance(d, dt.datetime) else d
        key = (d or dt.date.min, _ver_key(g("버전")), r)
        if pn not in latest or key > latest[pn][0]:
            latest[pn] = (key, r)

    ws = wb["품번별최신"]
    wb.remove(ws)
    ws = wb.create_sheet("품번별최신", wb.sheetnames.index("승인원대장") + 1)
    cols = [("품번", 15), ("품목이름", 34), ("업체 P/N", 20), ("적용모델", 12),
            ("공급업체", 15), ("최신버전", 9), ("승인구분", 14), ("승인일자", 12),
            ("최종판정", 9), ("발행No", 22), ("누적등록건수", 11), ("원본파일", 40)]
    style_header(ws, cols)
    rows = sorted(latest.items(), key=lambda kv: kv[0])
    for i, (pn, (_, r)) in enumerate(rows, 2):
        g = lambda n: ws_l.cell(row=r, column=COL[n]).value
        vals = [pn, g("품목이름"), g("업체 P/N"), g("적용모델"), g("공급업체"), g("버전"),
                g("승인구분"), g("승인일자"), g("최종판정"), g("발행No"), counts[pn], g("원본파일")]
        for ci, v in enumerate(vals, 1):
            ws.cell(row=i, column=ci, value=v)
        style_row(ws, i, len(cols), (2, 12))
        ws.cell(row=i, column=8).number_format = "yyyy-mm-dd"
        src = ws_l.cell(row=r, column=COL["원본파일"])
        if src.hyperlink:
            ws.cell(row=i, column=12).hyperlink = src.hyperlink.target
            ws.cell(row=i, column=12).font = F_LINK
    ws.auto_filter.ref = f"A1:L{max(len(rows) + 1, 1)}"
    ws.cell(row=len(rows) + 3, column=1,
            value="※ 이 시트는 실행 시마다 승인원대장 기준으로 자동 재생성됩니다 (직접 수정 금지).").font = \
        Font(name=FONT, size=9, italic=True, color="808080")


def rebuild_summary(wb):
    ws_l = wb["승인원대장"]
    rows = ledger_rows(ws_l)

    def distinct(name):
        s = OrderedDict()
        for r in rows:
            v = ws_l.cell(row=r, column=COL[name]).value
            if v not in (None, ""):
                s[v] = s.get(v, 0) + 1
        return [k for k, _ in sorted(s.items(), key=lambda kv: -kv[1])]

    old = wb["현황요약"]
    wb.remove(old)
    ws = wb.create_sheet("현황요약", 0)
    ws.sheet_view.showGridLines = False
    for c, w in zip("ABCDEF", (3, 22, 15, 15, 12, 3)):
        ws.column_dimensions[c].width = w

    L = "승인원대장"
    rng = lambda n: f"{L}!${get_column_letter(COL[n])}:${get_column_letter(COL[n])}"

    ws["B2"] = "승인원 관리 현황"
    ws["B2"].font = Font(name=FONT, size=16, bold=True, color="1F3864")
    ws["B3"] = "=\"기준: \"&TEXT(TODAY(),\"yyyy-mm-dd\")&\"  |  원본: 승인원대장 시트 (수식 자동집계)\""
    ws["B3"].font = Font(name=FONT, size=9, color="808080")

    kpi = [("총 등록건수", f"=COUNTA({rng('발행No')})-1"),
           ("관리 품번수", f"=COUNTA(품번별최신!$A:$A)-2"),
           ("합격", f"=COUNTIF({rng('최종판정')},\"합격\")"),
           ("불합격", f"=COUNTIF({rng('최종판정')},\"불합격\")"),
           ("확인필요(파싱)", f"=COUNTIF({rng('파싱상태')},\"확인필요*\")"),
           ("최근 30일 승인", f"=COUNTIFS({rng('승인일자')},\">=\"&(TODAY()-30))")]
    # KPI: 2행 x 3열 카드 (B~D)
    for i, (name, f) in enumerate(kpi):
        c = 2 + (i % 3)
        rr = 5 + (i // 3) * 2
        h = ws.cell(row=rr, column=c, value=name)
        h.font, h.fill, h.alignment, h.border = Font(name=FONT, size=9, bold=True, color="1F3864"), C_SUB, CENTER, BORDER
        v = ws.cell(row=rr + 1, column=c, value=f)
        v.font, v.alignment, v.border = Font(name=FONT, size=14, bold=True), CENTER, BORDER
        v.number_format = "#,##0"
    ws.column_dimensions["B"].width = 22

    def block(top, left, title, key, items, extra=None):
        hdr = [title, "건수", "합격", "불합격"]
        for j, h in enumerate(hdr):
            c = ws.cell(row=top, column=left + j, value=h)
            c.font, c.fill, c.alignment, c.border = F_HEAD, C_HEAD, CENTER, BORDER
        rr = top
        for it in items:
            rr += 1
            if isinstance(it, tuple):  # (표시명, COUNTIFS 조건)
                ws.cell(row=rr, column=left, value=it[0])
                cond = it[1]
            else:
                ws.cell(row=rr, column=left, value=it)
                cond = f"{rng(key)},{get_column_letter(left)}{rr}"
            ws.cell(row=rr, column=left + 1, value=f"=COUNTIFS({cond})")
            ws.cell(row=rr, column=left + 2, value=f"=COUNTIFS({cond},{rng('최종판정')},\"합격\")")
            ws.cell(row=rr, column=left + 3, value=f"=COUNTIFS({cond},{rng('최종판정')},\"불합격\")")
            for j in range(4):
                c = ws.cell(row=rr, column=left + j)
                c.font, c.border = F_BODY, BORDER
                c.alignment = LEFT if j == 0 else CENTER
        if not items:
            rr += 1
            ws.cell(row=rr, column=left, value="(데이터 없음)").font = F_BODY
        return rr

    # 업체별 / 승인구분별 / 모델별 / 월별
    top = 10
    end1 = block(top, 2, "공급업체별", "공급업체", distinct("공급업체"))
    top2 = end1 + 2
    end2 = block(top2, 2, "승인구분별", "승인구분", distinct("승인구분"))
    top3 = end2 + 2
    block(top3, 2, "적용모델별", "적용모델", distinct("적용모델"))

    # 월별 (승인일자 기준, 데이터 기간 전체)
    months = sorted({(d.year, d.month) for d in
                     (ws_l.cell(row=r, column=COL["승인일자"]).value for r in rows)
                     if isinstance(d, (dt.date, dt.datetime))})
    m_items = []
    for y, m in months:
        ny, nm = (y + 1, 1) if m == 12 else (y, m + 1)
        cond = (f"{rng('승인일자')},\">=\"&DATE({y},{m},1),{rng('승인일자')},\"<\"&DATE({ny},{nm},1)")
        m_items.append((f"{y}-{m:02d}", cond))
    ws.column_dimensions["G"].width = 16
    block(10, 7, "월별(승인일자)", None, m_items)
    for c in ("H", "I", "J"):
        ws.column_dimensions[c].width = 9


def sha256(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def backup():
    if not os.path.exists(XLSX):
        return
    os.makedirs(BACKUP_DIR, exist_ok=True)
    stamp = dt.datetime.now().strftime("%Y%m%d_%H%M%S")
    shutil.copy2(XLSX, os.path.join(BACKUP_DIR, f"승인원_관리대장_{stamp}.xlsx"))
    files = sorted(f for f in os.listdir(BACKUP_DIR) if f.endswith(".xlsx"))
    for f in files[:-BACKUP_KEEP]:
        os.remove(os.path.join(BACKUP_DIR, f))


def refresh_links(wb, extra=None):
    """원본파일 링크를 PC 절대경로로 재설정.
    상대경로 링크는 OneDrive 폴더에서 웹주소(onedrive.live.com)로 해석되어 열리지 않으므로,
    실행 시마다 보관 폴더를 스캔해 현재 위치 기준 절대경로로 맞춘다 (폴더 이동 시에도 자동 복구)."""
    files = dict(extra or {})
    for root, _, fs in os.walk(ARCHIVE):
        for f in fs:
            files.setdefault(f, os.path.join(root, f))
    changed = 0
    for sheet, col in (("승인원대장", COL["원본파일"]), ("품번별최신", 12)):
        if sheet not in wb.sheetnames:
            continue
        ws = wb[sheet]
        for r in range(2, ws.max_row + 1):
            c = ws.cell(row=r, column=col)
            path = files.get(str(c.value or ""))
            if not path:
                continue
            target = os.path.abspath(path)
            if not c.hyperlink or c.hyperlink.target != target:
                c.hyperlink = target
                c.font = F_LINK
                changed += 1
    return changed


def save_workbook(wb):
    wb.calculation.fullCalcOnLoad = True  # 엑셀 열 때 집계 수식 재계산
    backup()
    try:
        wb.save(XLSX)
        return True
    except PermissionError:
        print("\n[오류] 승인원_관리대장.xlsx 가 엑셀에서 열려 있어 저장할 수 없습니다.")
        print("       엑셀 파일을 닫고 다시 실행하세요. (PDF 는 이동/복사되지 않았습니다)")
        return False


LOG_COLS = [("처리일시", 17), ("구분", 9), ("파일명", 45), ("발행No", 22), ("내용", 70)]


def write_log(wb, entries):
    """처리로그 시트에 이번 실행 결과(등록/중복/실패/재분석) 누적"""
    if not entries:
        return
    if "처리로그" not in wb.sheetnames:
        style_header(wb.create_sheet("처리로그"), LOG_COLS)
    ws = wb["처리로그"]
    now = dt.datetime.now().replace(microsecond=0)
    r = ws.max_row
    while r > 1 and ws.cell(row=r, column=1).value is None:
        r -= 1
    for kind, fname, issue, msg in entries:
        r += 1
        for ci, v in enumerate((now, kind, fname, issue, msg), 1):
            ws.cell(row=r, column=ci, value=v)
        style_row(ws, r, len(LOG_COLS), (3, 5))
        ws.cell(row=r, column=1).number_format = "yyyy-mm-dd hh:mm"
        if kind in ("실패", "중복"):
            ws.cell(row=r, column=2).fill = C_BAD if kind == "실패" else C_NOTE
    ws.auto_filter.ref = f"A1:E{r}"


def explain(e):
    """자주 발생하는 오류를 한글 안내로 변환"""
    if isinstance(e, PermissionError):
        return ("승인원_관리대장.xlsx 파일에 접근할 수 없습니다.\n"
                "       ① 엑셀에서 파일이 열려 있으면 닫아주세요.\n"
                "       ② OneDrive 동기화 중(파란 화살표)이면 완료 후 다시 실행하세요.")
    if isinstance(e, FileNotFoundError):
        return f"파일을 찾을 수 없습니다: {getattr(e, 'filename', '')}"
    msg = str(e)
    if "No /Root object" in msg or "PDFSyntaxError" in type(e).__name__:
        return "PDF 파일이 손상되었거나 PDF 형식이 아닙니다."
    if "password" in msg.lower() or "encrypt" in msg.lower():
        return "암호가 걸린 PDF 입니다. 암호 해제 후 다시 넣어주세요."
    if "zip" in msg.lower() and "file" in msg.lower():
        return "승인원_관리대장.xlsx 파일이 손상되었습니다. 백업 폴더의 최신 파일로 교체하세요."
    return f"{type(e).__name__}: {msg}"


def load_ledger():
    try:
        return open_workbook()
    except Exception as e:
        print("\n[오류] " + explain(e))
        return None


def reanalyze():
    """보관 중인 PDF 를 개선된 인식 로직으로 다시 분석해 기존 행을 갱신 (No/등록일시/비고 유지)"""
    wb = load_ledger()
    if wb is None:
        return 2
    ws_l = wb["승인원대장"]
    files = {}
    for root, _, fs in os.walk(ARCHIVE):
        for f in fs:
            files.setdefault(f, os.path.join(root, f))
    items, hists, seen, logs, changed = {}, [], set(), [], 0
    for r in ledger_rows(ws_l):
        g = lambda n: ws_l.cell(row=r, column=COL[n]).value
        fname, issue = g("원본파일"), g("발행No")
        path = files.get(str(fname or ""))
        if not path:
            logs.append(("실패", fname, issue, "재분석: 보관 폴더에서 원본 PDF 를 찾지 못함 (기존 값 유지)"))
            continue
        try:
            rec = parse_pdf(path)
        except Exception as e:
            logs.append(("실패", fname, issue, "재분석: " + explain(e)))
            continue
        before = {k: g(k) for k in ("최종판정", "외관", "치수", "성능", "기타", "버전", "승인구분")}
        if "재등록" in str(g("파싱상태") or ""):
            rec["파싱상태"] += " / " + str(g("파싱상태")).split(" / ", 1)[-1]
        write_ledger_row(ws_l, r, rec)
        finish_ledger_row(ws_l, r, rec)
        diff = [f"{k}: {before[k] or '-'} → {rec.get(k) or '-'}" for k in before
                if (before[k] or "") != (rec.get(k) or "")]
        if diff:
            changed += 1
            logs.append(("재분석", fname, issue, "변경 " + ", ".join(diff)))
            print(f"[변경] {issue}  " + ", ".join(diff))
        base = {"발행No": rec.get("발행No"), "품번": rec.get("품번"), "버전": rec.get("버전")}
        items[r] = [dict(base, **it) for it in rec["_items"]]
        for h in rec["_hist"]:
            k = (rec.get("품번"), h["승인일자"], h["변경 후 버전"])
            if k not in seen:
                seen.add(k)
                hists.append(dict(base, **h))

    # 검토항목 / 승인이력 시트 재생성 (원본을 못 찾은 발행No 의 기존 행은 보존)
    reparsed = {ws_l.cell(row=r, column=COL["발행No"]).value for r in items}
    for title, cols, new_rows, wrap in (
            ("검토항목", ITEM_COLS, [x for v in items.values() for x in v], ("검토 및 시험사항",)),
            ("승인이력", HIST_COLS, hists, ("변경 전 내용", "변경 후 내용"))):
        old = wb[title]
        keep = [dict(zip([c for c, _ in cols], row)) for row in old.iter_rows(min_row=2, values_only=True)
                if row[0] and row[0] not in reparsed]
        pos = wb.sheetnames.index(title)
        wb.remove(old)
        ws = wb.create_sheet(title, pos)
        style_header(ws, cols)
        append_simple(ws, cols, keep + new_rows, wrap_names=wrap)

    rebuild_latest(wb)
    rebuild_summary(wb)
    refresh_links(wb)
    logs.append(("재분석", "", "", f"전체 {len(ledger_rows(ws_l))}건 재분석, 판정/버전 변경 {changed}건"))
    write_log(wb, logs)
    try:
        if not save_workbook(wb):
            return 2
    except Exception as e:
        print("\n[오류] " + explain(e))
        return 2
    for kind, fname, _, msg in logs:
        if kind == "실패":
            print(f"[실패] {fname} - {msg}")
    print(f"\n재분석 완료: 전체 {len(ledger_rows(ws_l))}건 / 판정·버전 변경 {changed}건 (상세: 처리로그 시트)")
    return 0


def main(argv):
    os.makedirs(INBOX, exist_ok=True)
    if argv and argv[0] in ("--재분석", "--reparse"):
        return reanalyze()
    if argv:
        sources, move = [os.path.abspath(p) for p in argv], False
    else:
        sources = [os.path.join(INBOX, f) for f in sorted(os.listdir(INBOX))
                   if f.lower().endswith(".pdf")]
        move = True
    sources = [p for p in sources if p.lower().endswith(".pdf") and os.path.isfile(p)]
    if not sources:
        print(f"[안내] 처리할 PDF 가 없습니다. '{os.path.relpath(INBOX, BASE)}' 폴더에 넣거나 PDF 를 드래그하세요.")
        if os.path.exists(XLSX):
            wb = load_ledger()
            if wb is None:
                return 2
            n = refresh_links(wb)
            if n:
                if not save_workbook(wb):
                    return 2
                print(f"[링크] 원본파일 링크 {n}건을 현재 PC 경로로 갱신했습니다.")
        return 0

    wb = load_ledger()
    if wb is None:
        print("       (PDF 는 수신함에 그대로 남아 있습니다)")
        return 2
    ws_l = wb["승인원대장"]
    known_hash = {str(ws_l.cell(row=r, column=COL["파일해시"]).value): r for r in ledger_rows(ws_l)}
    known_issue = {(ws_l.cell(row=r, column=COL["발행No"]).value,
                    ws_l.cell(row=r, column=COL["버전"]).value): r for r in ledger_rows(ws_l)}

    ws_h = wb["승인이력"]
    known_hist = set()
    for r in range(2, ws_h.max_row + 1):
        d = ws_h.cell(row=r, column=4).value
        d = d.date() if isinstance(d, dt.datetime) else d
        known_hist.add((ws_h.cell(row=r, column=2).value, d, ws_h.cell(row=r, column=6).value))

    plan, failed, skipped = [], [], []
    for src in sources:
        name = os.path.basename(src)
        sha = sha256(src)
        if sha[:12] in known_hash:
            skipped.append((src, f"이미 등록됨 (대장 {known_hash[sha[:12]] - 1}번)"))
            continue
        try:
            rec = parse_pdf(src)
        except Exception as e:  # 손상/암호화 PDF 등
            failed.append((src, explain(e)))
            continue
        if not rec.get("품번") and not rec.get("발행No"):
            failed.append((src, "승인서 양식을 인식하지 못함 (품번/발행No 없음)"))
            continue
        dup = known_issue.get((rec.get("발행No"), rec.get("버전")))
        if dup:
            rec["파싱상태"] = (rec["파싱상태"] if rec["파싱상태"] != "정상" else "확인필요") + \
                f" / 동일 발행No·버전 재등록(대장 {dup - 1}번)"
        rel = archive_name(rec, src)
        r = append_ledger(ws_l, rec, rel, sha)
        known_hash[sha[:12]] = r
        known_issue[(rec.get("발행No"), rec.get("버전"))] = r
        base = {"발행No": rec.get("발행No"), "품번": rec.get("품번"), "버전": rec.get("버전")}
        append_simple(wb["검토항목"], ITEM_COLS, [dict(base, **it) for it in rec["_items"]],
                      wrap_names=("검토 및 시험사항",))
        # 승인원마다 과거 이력이 반복 수록되므로 (품번, 승인일자, 변경 후 버전) 기준 신규분만 추가
        new_hist = []
        for h in rec["_hist"]:
            k = (rec.get("품번"), h["승인일자"], h["변경 후 버전"])
            if k not in known_hist:
                known_hist.add(k)
                new_hist.append(dict(base, **h))
        append_simple(wb["승인이력"], HIST_COLS, new_hist,
                      wrap_names=("변경 전 내용", "변경 후 내용"))
        plan.append((src, rel, rec))
        print(f"[등록] {name}\n       -> {rec.get('발행No')} | {rec.get('품번')} | {rec.get('승인구분')} "
              f"{rec.get('버전')} | {rec.get('승인일자')} | {rec.get('최종판정')} | {rec['파싱상태']}")

    logs = [("등록", os.path.basename(src), rec.get("발행No"),
             f"{rec.get('품번')} {rec.get('승인구분')} {rec.get('버전')} / {rec.get('최종판정')} / {rec['파싱상태']}")
            for src, _, rec in plan]
    logs += [("중복", os.path.basename(src), "", why) for src, why in skipped]
    logs += [("실패", os.path.basename(src), "", why) for src, why in failed]
    if plan:
        rebuild_latest(wb)
        rebuild_summary(wb)
        refresh_links(wb, {os.path.basename(rel): os.path.join(BASE, rel) for _, rel, _ in plan})
    write_log(wb, logs)
    if logs:
        if not save_workbook(wb):
            return 2
        for src, rel, _ in plan:
            dst = os.path.join(BASE, rel)
            os.makedirs(os.path.dirname(dst), exist_ok=True)
            (shutil.move if move else shutil.copy2)(src, dst)

    if move:  # 수신함 정리: 중복 -> 중복/, 실패 -> 오류/
        for group, folder in ((skipped, DUP_DIR), (failed, ERROR_DIR)):
            for src, _ in group:
                os.makedirs(folder, exist_ok=True)
                shutil.move(src, os.path.join(folder, os.path.basename(src)))
    for src, why in skipped:
        print(f"[중복] {os.path.basename(src)} - {why} → 승인원_수신함\\중복 폴더로 이동" if move else
              f"[중복] {os.path.basename(src)} - {why}")
    for src, why in failed:
        print(f"[실패] {os.path.basename(src)} - {why}" + (" → 승인원_수신함\\오류 폴더로 이동" if move else ""))
    print(f"\n완료: 등록 {len(plan)}건 / 중복 {len(skipped)}건 / 실패 {len(failed)}건 (상세: 처리로그 시트)")
    if plan:
        print(f"대장: {XLSX}")
    return 0 if not failed else 1


if __name__ == "__main__":
    try:
        sys.exit(main(sys.argv[1:]))
    except Exception as e:  # 예상 못한 오류: 한글 안내 + 상세 로그 파일
        import traceback
        with open(os.path.join(BASE, "오류로그.txt"), "a", encoding="utf-8") as f:
            f.write(f"\n[{dt.datetime.now():%Y-%m-%d %H:%M:%S}]\n{traceback.format_exc()}")
        print("\n[오류] " + explain(e))
        print("       상세 내용은 오류로그.txt 에 저장했습니다. 화면 캡처와 함께 전달해 주세요.")
        sys.exit(3)
